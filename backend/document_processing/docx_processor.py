from pathlib import Path
import docx
from backend.document_processing.base_processor import DocumentProcessor, ProcessedContent

class DOCXProcessor(DocumentProcessor):
    def extract(self, file_path: Path) -> ProcessedContent:
        try:
            doc = docx.Document(str(file_path))
        except Exception as exc:
            return ProcessedContent(
                text="",
                page_count=1,
                word_count=0,
                language="en",
                metadata={"extraction_warning": f"DOCX package could not be parsed safely: {type(exc).__name__}"},
                has_ocr=False
            )
        text_parts = []

        for p in doc.paragraphs:
            if p.text.strip():
                text_parts.append(p.text)

        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_parts.append(" | ".join(row_text))

        full_text = "\n".join(text_parts).strip()
        word_count = len(full_text.split()) if full_text else 0

        metadata = {}
        try:
            core_props = doc.core_properties
            metadata = {
                "title": core_props.title or "",
                "author": core_props.author or "",
                "subject": core_props.subject or "",
                "created": str(core_props.created) if core_props.created else "",
                "modified": str(core_props.modified) if core_props.modified else ""
            }
        except Exception:
            pass

        lang = "en"
        if full_text:
            try:
                import langdetect
                lang = langdetect.detect(full_text[:1000])
            except Exception:
                lang = "en"

        return ProcessedContent(
            text=full_text,
            page_count=1,
            word_count=word_count,
            language=lang,
            metadata=metadata,
            has_ocr=False
        )
